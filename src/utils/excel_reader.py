from pathlib import Path

from openpyxl import load_workbook


DEFAULT_FILE = Path(__file__).resolve().parents[2] / "test_cases" / "login_test_cases.xlsx"


def load_login_test_cases(file_path=DEFAULT_FILE):
    workbook = load_workbook(file_path, data_only=True, read_only=True)
    worksheet = workbook["Login"]
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value).strip() for value in next(rows)]
    test_cases = []

    for row in rows:
        test_case = dict(zip(headers, row))
        if not test_case.get("test_id"):
            continue
        if str(test_case.get("enabled", "yes")).strip().lower() not in {"yes", "true", "1"}:
            continue
        test_cases.append({key: value if value is not None else "" for key, value in test_case.items()})

    workbook.close()
    return test_cases

